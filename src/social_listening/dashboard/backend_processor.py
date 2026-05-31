from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from social_listening.dashboard.repository import PostgresDashboardRepository, sql_json, sql_str
from social_listening.text_utils import normalize_text


SCREEN_SHEET_MAP = {
    "overview": "Overview",
    "listen": "Listen",
    "brand": "Brand Health",
    "reputation": "Reputation",
}


@dataclass(frozen=True)
class DashboardRequirement:
    screen: str
    dashboard: str
    purpose: str
    input_data: str
    formula: str
    output_metrics: str
    action_owner: str
    data_gap: str = ""
    priority: str = ""
    formula_readiness: str = ""
    ba_dev_note: str = ""


class DashboardBackendProcessor:
    """Build the dashboard-ready snapshot layer from PostgreSQL and Excel formula mapping.

    The web dashboard should render this snapshot instead of recalculating business
    logic in the browser. The mapping workbook is treated as the contract.
    """

    def __init__(
        self,
        repository: PostgresDashboardRepository,
        mapping_file: Path,
    ) -> None:
        self.repository = repository
        self.mapping_file = Path(mapping_file)
        self.requirements = load_dashboard_requirements(self.mapping_file)

    def build_snapshot(self) -> dict[str, Any]:
        payload = self.repository.get_dashboard_payload()
        if not payload:
            raise RuntimeError("No dashboard payload returned from PostgreSQL")

        screens = json.loads(json.dumps(payload.get("screens") or {}))
        context = DashboardContext(
            payload,
            mentions=self.repository._mentions(),
            reviews=self.repository._reviews(),
            menu_items=self.repository._menu_items(),
        )
        self._rewrite_overview_top_cards(screens.setdefault("overview", {}), context)
        self._complete_overview(screens.setdefault("overview", {}), context)
        self._complete_reputation(screens.setdefault("reputation", {}), context)
        self._attach_requirement_metadata(screens)
        validation = self._validate_contract(screens)
        if validation["missing"]:
            missing = "; ".join(f"{item['screen']}:{item['dashboard']}" for item in validation["missing"])
            raise RuntimeError(f"Dashboard backend contract still missing modules: {missing}")

        overview_screen = screens.get("overview") or {}
        snapshot = {
            "backend": {
                "name": "dashboard_backend_processor",
                "version": 1,
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "mapping_file": str(self.mapping_file),
                "validation": validation,
            },
            "overview": {
                "headline": overview_screen.get("headline") or payload.get("overview", {}).get("headline") or "",
                "top_cards": overview_screen.get("topCards") or payload.get("overview", {}).get("top_cards") or [],
            },
            "screens": screens,
            "requirements": {
                key: [asdict(item) for item in values]
                for key, values in self.requirements.items()
            },
        }
        return snapshot

    def _rewrite_overview_top_cards(self, screen: dict[str, Any], context: "DashboardContext") -> None:
        cards = list(screen.get("topCards") or [])
        opportunity = context.top_food_opportunity()
        if not opportunity:
            return
        rewritten = []
        replaced = False
        for card in cards:
            if str(card.get("tag") or "").strip().upper() == "CƠ HỘI TĂNG TRƯỞNG":
                rewritten.append({
                    "tag": "CƠ HỘI TĂNG TRƯỞNG",
                    "title": opportunity["title"],
                    "value": "Sẵn sàng" if opportunity["count"] >= 5 else "Theo dõi",
                    "desc": f"{opportunity['count']} bình luận/review nhắc tới {opportunity['label']} trong dữ liệu hiện tại.",
                    "severity": "medium" if opportunity["count"] >= 5 else "low",
                    "owner": "Marketing",
                    "guardrail": "Sẵn sàng chạy campaign" if opportunity["count"] >= 5 else "Theo dõi thêm",
                    "evidence_query": "Trend Opportunity Snapshot",
                    "evidence_kind": "food_opportunity",
                    "evidence_platform": "",
                    "metadata": {
                        "formula": "Top food opportunity = most-mentioned menu/dish alias across mention/review text",
                        "matched_aliases": opportunity["aliases"],
                    },
                })
                replaced = True
            else:
                rewritten.append(card)
        if not replaced:
            rewritten.insert(2, {
                "tag": "CƠ HỘI TĂNG TRƯỞNG",
                "title": opportunity["title"],
                "value": "Sẵn sàng" if opportunity["count"] >= 5 else "Theo dõi",
                "desc": f"{opportunity['count']} bình luận/review nhắc tới {opportunity['label']} trong dữ liệu hiện tại.",
                "severity": "medium" if opportunity["count"] >= 5 else "low",
                "owner": "Marketing",
                "guardrail": "Sẵn sàng chạy campaign" if opportunity["count"] >= 5 else "Theo dõi thêm",
                "evidence_query": "Trend Opportunity Snapshot",
                "evidence_kind": "food_opportunity",
            })
        screen["topCards"] = rewritten

    def save_snapshot(self, snapshot: dict[str, Any]) -> None:
        brand_id = self.repository._brand_id()
        if not brand_id:
            raise RuntimeError(f"Brand slug not found: {self.repository.brand_slug}")
        sql = f"""
        INSERT INTO {self.repository.schema}.dashboard_snapshots (
          brand_id, snapshot_date, scope_type, scope_key, payload
        )
        VALUES (
          {sql_str(brand_id)},
          CURRENT_DATE,
          'brand',
          'all',
          {sql_json(snapshot)}::jsonb
        )
        ON CONFLICT (brand_id, snapshot_date, scope_type, scope_key)
        DO UPDATE SET payload = EXCLUDED.payload, created_at = NOW();
        """
        self.repository._execute_sql(sql)

    def _complete_overview(self, screen: dict[str, Any], context: "DashboardContext") -> None:
        modules = list(screen.get("modules") or [])
        present = {str(module.get("title") or "") for module in modules if isinstance(module, dict)}
        builders = {
            "Category Pulse": self._build_category_pulse,
            "Trend Opportunity Snapshot": self._build_trend_opportunity_snapshot,
            # "Priority Action Detail" removed - it's a FEED TABLE, not a dashboard module
        }
        for title, builder in builders.items():
            if title not in present:
                modules.append(builder(context, screen_key="overview"))
        screen["modules"] = order_modules("overview", modules, self.requirements)

    def _complete_reputation(self, screen: dict[str, Any], context: "DashboardContext") -> None:
        modules = list(screen.get("modules") or [])
        present = {str(module.get("title") or "") for module in modules if isinstance(module, dict)}
        builders = {
            "Branch x Issue Heatmap": self._build_branch_issue_heatmap,
            # "Priority Action Detail" removed - it's a FEED TABLE, not a dashboard module
        }
        for title, builder in builders.items():
            if title not in present:
                modules.append(builder(context, screen_key="reputation"))
        screen["modules"] = order_modules("reputation", modules, self.requirements)

    def _attach_requirement_metadata(self, screens: dict[str, Any]) -> None:
        for screen_key, screen in screens.items():
            req_by_title = {
                item.dashboard: item
                for item in self.requirements.get(screen_key, [])
            }
            for module in screen.get("modules") or []:
                if not isinstance(module, dict):
                    continue
                requirement = req_by_title.get(str(module.get("title") or ""))
                if requirement:
                    module["formula_ref"] = asdict(requirement)

    def _validate_contract(self, screens: dict[str, Any]) -> dict[str, Any]:
        missing: list[dict[str, str]] = []
        present_by_screen: dict[str, list[str]] = {}
        for screen_key, requirements in self.requirements.items():
            module_titles = [
                str(module.get("title") or "")
                for module in (screens.get(screen_key, {}).get("modules") or [])
                if isinstance(module, dict)
            ]
            present_by_screen[screen_key] = module_titles
            for requirement in requirements:
                if requirement.dashboard not in module_titles:
                    missing.append({"screen": screen_key, "dashboard": requirement.dashboard})
        return {"missing": missing, "present_by_screen": present_by_screen}

    def _build_category_pulse(self, context: "DashboardContext", screen_key: str) -> dict[str, Any]:
        rows = []
        for topic, stats in context.topic_stats().items():
            total = max(1, stats["positive"] + stats["negative"] + stats["neutral"])
            sentiment_health = ((stats["positive"] - stats["negative"]) / total) * 100
            signal_quality = average(stats["confidence"]) if stats["confidence"] else 70
            health = clamp(round(sentiment_health * 0.60 + signal_quality * 0.40, 2), 0, 100)
            status = "Healthy" if health >= 75 else ("Watch" if health >= 50 else "Critical")
            rows.append([humanize_topic(topic), int(health), status])
        rows.sort(key=lambda row: row[1])
        return {
            "title": "Category Pulse",
            "takeaway": "Đo health theo category/topic để biết vùng sản phẩm hoặc trải nghiệm nào cần xử lý trước.",
            "chart": "hbar",
            "data": [[row[0], row[1], "high" if row[2] == "Critical" else ("medium" if row[2] == "Watch" else "low")] for row in rows[:8]],
            "analysis": [
                "Health score dùng sentiment health và signal quality theo mapping.",
                "Topic có health thấp nhưng evidence yếu sẽ không được đẩy thành kết luận mạnh.",
            ],
            "action": {"guardrail": "Quality-adjusted category health", "owner": "Product team", "cta": "Focus weak categories"},
            "evidence_query": "Category Pulse",
        }

    def _build_trend_opportunity_snapshot(self, context: "DashboardContext", screen_key: str) -> dict[str, Any]:
        rows = []
        for topic, stats in context.topic_stats().items():
            demand_velocity = clamp(stats["count"] * 8 + stats["positive"] * 6, 0, 100)
            brand_fit = 65 if topic in {"food_quality", "menu_variety", "price_value"} else 55
            audience_fit = 60 if stats["count"] >= 2 else 45
            risk_inverse = 100 - clamp(stats["negative"] * 18, 0, 100)
            evidence_confidence = average(stats["confidence"]) if stats["confidence"] else 65
            readiness = round(
                demand_velocity * 0.30
                + brand_fit * 0.25
                + audience_fit * 0.20
                + risk_inverse * 0.15
                + evidence_confidence * 0.10,
                2,
            )
            rows.append([humanize_topic(topic), int(readiness), class_for_score(readiness)])
        rows.sort(key=lambda row: row[1], reverse=True)
        return {
            "title": "Trend Opportunity Snapshot",
            "takeaway": "Chỉ đề xuất test 7 ngày khi trend có demand velocity, brand fit và evidence confidence đủ.",
            "chart": "hbar",
            "data": [[row[0], row[1], "low" if row[2].startswith("Ready") else "medium"] for row in rows[:6]],
            "analysis": [
                "Readiness theo công thức demand_velocity, brand_fit, audience_fit, risk_inverse và confidence.",
                "Guardrail của workbook: chỉ suggest 7-day test, không full rollout.",
            ],
            "action": {"guardrail": "7-day test only", "owner": "Marketing team", "cta": "Launch trend campaigns"},
            "evidence_query": "Trend Opportunity Snapshot",
        }

    def _build_priority_action_detail(self, context: "DashboardContext", screen_key: str) -> dict[str, Any]:
        rows = []
        for item in context.action_feed[:10]:
            severity = severity_score(item.get("severity"))
            confidence = 75 if item.get("summary") else 60
            business_impact = 75 if severity >= 75 else 55
            urgency = 80 if severity >= 75 else 55
            priority = priority_score(severity, business_impact, urgency, confidence)
            sla = "24h" if priority >= 75 else ("48h" if priority >= 50 else "72h")
            rows.append([
                item.get("title") or "",
                f"Priority {int(priority)}",
                f"{item.get('owner') or 'Owner'} · SLA {sla} · {item.get('status') or 'Need more data'}",
            ])
        return {
            "title": "Priority Action Detail",
            "takeaway": "Action queue được sinh từ priority score, SLA và evidence guardrail, không để UI tự suy luận.",
            "chart": "list",
            "data": rows,
            "analysis": [
                "Priority score = severity*0.35 + business_impact*0.25 + urgency*0.20 + confidence*0.20.",
                "Confidence hoặc evidence thấp sẽ giữ trạng thái Need more data thay vì auto-assign mạnh.",
            ],
            "action": {"guardrail": "Evidence-backed SLA", "owner": "All teams", "cta": "Follow timeline execution"},
            "evidence_query": "Priority Action Detail",
        }

    def _build_branch_issue_heatmap(self, context: "DashboardContext", screen_key: str) -> dict[str, Any]:
        topics = context.top_topics(limit=5)
        rows = [["Branch", *[humanize_topic(topic) for topic in topics]]]
        for branch in context.branch_intelligence[:8]:
            branch_topics = set(branch.get("top_topics") or [])
            row = [branch.get("branch_name") or branch.get("branch_slug") or ""]
            for topic in topics:
                value = "High" if topic in branch_topics and int(branch.get("risk_score") or 0) >= 75 else (
                    "Watch" if topic in branch_topics else "OK"
                )
                row.append(value)
            rows.append(row)
        return {
            "title": "Branch x Issue Heatmap",
            "takeaway": "Heatmap giúp tách issue toàn hệ thống khỏi pain cục bộ theo từng chi nhánh.",
            "chart": "table",
            "data": rows,
            "analysis": [
                "Brand-wide issue khi affected_branch_ratio >= 0.60 hoặc affected_branch_count >= 3.",
                "Branch-specific issue cần branch owner, không nên biến thành brand-wide reaction.",
            ],
            "action": {"guardrail": "Do not over-generalize local pain", "owner": "CS + Branch managers", "cta": "Address weak cells"},
            "evidence_query": "Branch x Issue Heatmap",
        }


class DashboardContext:
    def __init__(
        self,
        payload: dict[str, Any],
        mentions: list[dict[str, Any]] | None = None,
        reviews: list[dict[str, Any]] | None = None,
        menu_items: list[dict[str, Any]] | None = None,
    ) -> None:
        self.payload = payload
        self.platform_summary = payload.get("platform_summary") or []
        self.branch_intelligence = payload.get("branch_intelligence") or []
        self.evidence_cards = payload.get("evidence_cards") or []
        self.action_feed = payload.get("mentions_feed") or []
        self.mentions = mentions or []
        self.reviews = reviews or []
        self.menu_items = menu_items or []

    def topic_stats(self) -> dict[str, dict[str, Any]]:
        stats: dict[str, dict[str, Any]] = {}
        for branch in self.branch_intelligence:
            for topic in branch.get("top_topics") or ["unknown"]:
                row = stats.setdefault(str(topic), {"count": 0, "positive": 0, "negative": 0, "neutral": 0, "confidence": []})
                row["count"] += 1
                row["positive"] += int(branch.get("positive_count") or 0)
                row["negative"] += int(branch.get("negative_count") or 0)
                row["neutral"] += max(0, int(branch.get("review_count") or 0) - int(branch.get("positive_count") or 0) - int(branch.get("negative_count") or 0))
                row["confidence"].append(75)
        for card in self.evidence_cards:
            topic = infer_topic_from_text(str(card.get("evidence_quote") or ""))
            row = stats.setdefault(topic, {"count": 0, "positive": 0, "negative": 0, "neutral": 0, "confidence": []})
            row["count"] += 1
            sentiment = str(card.get("sentiment_label") or "neutral")
            if sentiment == "positive":
                row["positive"] += 1
            elif sentiment == "negative":
                row["negative"] += 1
            else:
                row["neutral"] += 1
            row["confidence"].append(float(card.get("confidence_score") or 0.7) * 100)
        return stats

    def top_topics(self, limit: int = 5) -> list[str]:
        stats = self.topic_stats()
        return [
            topic
            for topic, _ in sorted(stats.items(), key=lambda item: item[1]["count"], reverse=True)[:limit]
        ]

    def top_food_opportunity(self) -> dict[str, Any] | None:
        families = build_food_families(self.menu_items)
        if not families:
            return None
        texts = []
        for mention in self.mentions:
            text = str(mention.get("content_text") or "").strip()
            if text:
                texts.append(text)
        for review in self.reviews:
            text = str(review.get("review_text") or "").strip()
            if text:
                texts.append(text)
        if not texts:
            return None
        scores: dict[str, dict[str, Any]] = {
            key: {"count": 0, "label": value["label"], "aliases": value["aliases"]}
            for key, value in families.items()
        }
        for text in texts:
            normalized = searchable_text(text)
            matched_in_text: set[str] = set()
            for key, family in families.items():
                if any(alias and alias in normalized for alias in family["aliases"]):
                    matched_in_text.add(key)
            for key in matched_in_text:
                scores[key]["count"] += 1
        best = max(scores.values(), key=lambda item: item["count"], default=None)
        if not best or int(best["count"]) <= 0:
            return None
        return {
            "title": best["label"],
            "label": best["label"].lower(),
            "count": int(best["count"]),
            "aliases": best["aliases"][:8],
        }


def load_dashboard_requirements(mapping_file: Path) -> dict[str, list[DashboardRequirement]]:
    wb = load_workbook(mapping_file, data_only=True)
    requirements: dict[str, list[DashboardRequirement]] = {}
    gap_notes = load_gap_notes(wb)
    for screen_key, sheet_name in SCREEN_SHEET_MAP.items():
        ws = wb[sheet_name]
        rows: list[DashboardRequirement] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            dashboard = clean(row[0])
            if not dashboard:
                continue
            gap = gap_notes.get(dashboard, {})
            rows.append(
                DashboardRequirement(
                    screen=screen_key,
                    dashboard=dashboard,
                    purpose=clean(row[1]),
                    input_data=clean(row[2] or (row[6] if len(row) > 6 else "")),
                    formula=clean(row[3]),
                    output_metrics=clean(row[4]),
                    action_owner=clean(row[5]),
                    data_gap=gap.get("missing", ""),
                    priority=gap.get("priority", ""),
                    formula_readiness=gap.get("readiness", ""),
                    ba_dev_note=gap.get("note", ""),
                )
            )
        requirements[screen_key] = rows
    return requirements


def load_gap_notes(wb: Any) -> dict[str, dict[str, str]]:
    if "Gap Analysis" not in wb.sheetnames:
        return {}
    ws = wb["Gap Analysis"]
    notes = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        dashboard = clean(row[0])
        if not dashboard:
            continue
        notes[dashboard] = {
            "missing": clean(row[3]),
            "priority": clean(row[5]),
            "readiness": clean(row[9]),
            "note": clean(row[10]),
        }
    return notes


def order_modules(screen_key: str, modules: list[dict[str, Any]], requirements: dict[str, list[DashboardRequirement]]) -> list[dict[str, Any]]:
    order = {item.dashboard: index for index, item in enumerate(requirements.get(screen_key, []))}
    return sorted(modules, key=lambda module: order.get(str(module.get("title") or ""), 10_000))


def priority_score(severity: float, business_impact: float, urgency: float, confidence: float) -> float:
    return round(severity * 0.35 + business_impact * 0.25 + urgency * 0.20 + confidence * 0.20, 2)


def severity_score(value: object) -> float:
    text = str(value or "").lower()
    if text == "high":
        return 100
    if text == "medium":
        return 60
    if text == "low":
        return 30
    try:
        return float(value)
    except (TypeError, ValueError):
        return 30


def class_for_score(value: float) -> str:
    if value >= 75:
        return "Ready for 7-day test"
    if value >= 50:
        return "Watch / Caution"
    return "Not ready"


def infer_topic_from_text(text: str) -> str:
    lowered = text.casefold()
    if any(word in lowered for word in ("giao", "ship", "delivery", "nguội", "chậm")):
        return "delivery"
    if any(word in lowered for word in ("giá", "đắt", "không đáng", "value")):
        return "price_value"
    if any(word in lowered for word in ("nhân viên", "phục vụ", "service")):
        return "staff_service"
    if any(word in lowered for word in ("mì", "bò", "sủi cảo", "ngon", "mặn", "nhạt")):
        return "food_quality"
    return "general"


def build_food_families(menu_items: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    families: dict[str, dict[str, Any]] = {}
    for item in menu_items:
        item_name = str(item.get("item_name") or "").strip()
        if not item_name:
            continue
        family_key, label, aliases = classify_menu_family(item_name)
        if not family_key:
            continue
        row = families.setdefault(family_key, {"label": label, "aliases": set()})
        row["aliases"].update(aliases)
    return {
        key: {"label": value["label"], "aliases": sorted(value["aliases"], key=len, reverse=True)}
        for key, value in families.items()
    }


def classify_menu_family(item_name: str) -> tuple[str, str, set[str]]:
    text = searchable_text(item_name)
    aliases: set[str] = set()
    if any(token in text for token in ("mi suon bo", "mi bo", "beef noodle", "bo dai loan")):
        aliases.update({"mi bo", "mi suon bo", "beef noodle", "beef noodle soup", "taiwanese beef noodle", "bo dai loan"})
        return "beef_noodle", "Mì bò Đài Loan", aliases
    if any(token in text for token in ("sui cao", "dumpling", "dumplings")):
        aliases.update({"sui cao", "dumpling", "dumplings", "wonton"})
        return "dumpling", "Sủi cảo", aliases
    if any(token in text for token in ("com chien", "rice", "fried rice")):
        aliases.update({"com chien", "rice dish", "fried rice", "rice"})
        return "rice", "Cơm / cơm chiên", aliases
    if any(token in text for token in ("ga gion", "ga quay", "chicken")):
        aliases.update({"ga gion", "ga quay", "chicken", "crispy chicken"})
        return "chicken", "Gà giòn / gà quay", aliases
    if any(token in text for token in ("pao", "bao", "banh bao")):
        aliases.update({"pao", "bao", "banh bao"})
        return "pao", "Pao / bánh bao kẹp", aliases
    if any(token in text for token in ("tra", "hong tra", "7up", "sting", "fanta")):
        aliases.update({"tra", "hong tra", "tea", "drink", "drinks"})
        return "drink", "Đồ uống", aliases
    return "", "", set()


def searchable_text(value: str) -> str:
    text = normalize_text(value)
    replacements = {
        "đ": "d",
        "Đ": "d",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    import unicodedata

    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return " ".join(text.split())


def humanize_topic(value: object) -> str:
    labels = {
        "delivery": "Delivery friction",
        "price_value": "Price / value",
        "staff_service": "Staff service",
        "service_speed": "Service speed",
        "food_quality": "Food quality",
        "menu_variety": "Menu variety",
        "location": "Location / branch",
        "general": "General signal",
        "unknown": "Unknown",
    }
    text = str(value or "unknown")
    return labels.get(text, text.replace("_", " ").title())


def average(values: list[float]) -> float:
    return sum(values) / len(values) if values else 0


def clamp(value: float, low: float, high: float) -> float:
    return max(low, min(high, value))


def clean(value: object) -> str:
    return str(value or "").strip()
