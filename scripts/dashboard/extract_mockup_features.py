"""
Extract all features from mockup HTML to create comprehensive Excel specification
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import json
import re

def extract_data_from_html(html_path):
    """Extract JavaScript DATA object from HTML file"""
    with open(html_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Find the DATA object
    match = re.search(r'const DATA = ({.*?});', content, re.DOTALL)
    if not match:
        raise ValueError("Could not find DATA object in HTML")

    data_str = match.group(1)
    # Parse the JSON-like structure
    data = json.loads(data_str)
    return data

def create_feature_excel(data):
    """Create comprehensive Excel with all mockup features"""
    wb = Workbook()

    # Remove default sheet
    wb.remove(wb.active)

    # Define colors
    header_fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    screen_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    screen_font = Font(bold=True, color="FFFFFF", size=13)
    module_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # Border styles
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    screens = {
        'overview': 'Overview',
        'listen': 'Listen',
        'brand': 'Brand Health',
        'reputation': 'Reputation',
        'competitor': 'Competitor Radar'
    }

    # Create summary sheet first
    summary_ws = wb.create_sheet("📊 Summary", 0)

    # Summary headers
    summary_ws.append(["Screen", "Modules Count", "Chart Types", "Description"])
    for col in range(1, 5):
        cell = summary_ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    total_modules = 0

    # Process each screen
    for screen_key, screen_name in screens.items():
        if screen_key not in data:
            continue

        screen_data = data[screen_key]
        modules = screen_data.get('modules', [])

        # Count chart types
        chart_types = set()
        for module in modules:
            chart_types.add(module.get('chart', 'unknown'))

        # Add to summary
        summary_ws.append([
            screen_name,
            len(modules),
            ', '.join(sorted(chart_types)),
            screen_data.get('subtitle', '')
        ])

        total_modules += len(modules)

        # Create detail sheet for this screen
        ws = wb.create_sheet(f"🎯 {screen_name}")

        # Screen header
        ws.merge_cells('A1:G1')
        screen_header = ws['A1']
        screen_header.value = f"{screen_name} - {screen_data.get('subtitle', '')}"
        screen_header.fill = screen_fill
        screen_header.font = screen_font
        screen_header.alignment = Alignment(horizontal="center", vertical="center")

        # Headline
        ws.merge_cells('A2:G2')
        headline_cell = ws['A2']
        headline_cell.value = f"💡 Headline: {screen_data.get('headline', '')}"
        headline_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        headline_cell.font = Font(italic=True, size=10)
        ws.row_dimensions[2].height = 40

        # Column headers
        ws.append([])  # Empty row
        headers = ["#", "Module Name", "Chart Type", "Takeaway", "Analysis Points", "Action Owner", "Guardrail"]
        ws.append(headers)

        header_row = ws.max_row
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Module data
        for idx, module in enumerate(modules, 1):
            # Extract analysis points
            analysis_list = module.get('analysis', [])
            analysis_text = '\n'.join([f"• {point}" for point in analysis_list])

            action = module.get('action', {})

            row_data = [
                idx,
                module.get('title', ''),
                module.get('chart', ''),
                module.get('takeaway', ''),
                analysis_text,
                action.get('owner', ''),
                action.get('guardrail', '')
            ]

            ws.append(row_data)
            current_row = ws.max_row

            # Apply module styling
            for col_num in range(1, 8):
                cell = ws.cell(row=current_row, column=col_num)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

                # Module name - bold
                if col_num == 2:
                    cell.font = Font(bold=True)
                    cell.fill = module_fill

            # Set row height for better readability
            ws.row_dimensions[current_row].height = max(50, len(analysis_list) * 15)

        # Column widths
        ws.column_dimensions['A'].width = 5   # #
        ws.column_dimensions['B'].width = 35  # Module Name
        ws.column_dimensions['C'].width = 15  # Chart Type
        ws.column_dimensions['D'].width = 50  # Takeaway
        ws.column_dimensions['E'].width = 60  # Analysis
        ws.column_dimensions['F'].width = 20  # Owner
        ws.column_dimensions['G'].width = 25  # Guardrail

        # Add Top Cards section at bottom
        ws.append([])
        ws.append([])
        ws.merge_cells(f'A{ws.max_row}:G{ws.max_row}')
        topcard_header = ws.cell(row=ws.max_row, column=1)
        topcard_header.value = "📌 Top Cards (Quick Metrics)"
        topcard_header.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        topcard_header.font = Font(bold=True, color="FFFFFF", size=12)
        topcard_header.alignment = Alignment(horizontal="left", vertical="center")

        # Top cards headers
        ws.append(["Tag", "Title", "Value", "Description", "Owner", "Guardrail", "Severity"])
        topcard_header_row = ws.max_row
        for col_num in range(1, 8):
            cell = ws.cell(row=topcard_header_row, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Top cards data
        for card in screen_data.get('topCards', []):
            ws.append([
                card.get('tag', ''),
                card.get('title', ''),
                card.get('value', ''),
                card.get('desc', ''),
                card.get('owner', ''),
                card.get('guardrail', ''),
                card.get('sev', '')
            ])
            current_row = ws.max_row
            for col_num in range(1, 8):
                cell = ws.cell(row=current_row, column=col_num)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            # Color code severity
            sev_cell = ws.cell(row=current_row, column=7)
            sev = card.get('sev', '').lower()
            if sev == 'high':
                sev_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            elif sev == 'medium':
                sev_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            elif sev == 'low':
                sev_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    # Add totals to summary
    summary_ws.append([])
    summary_ws.append(["TOTAL", total_modules, "", ""])
    total_row = summary_ws.max_row
    for col in range(1, 3):
        cell = summary_ws.cell(row=total_row, column=col)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")

    # Adjust summary column widths
    summary_ws.column_dimensions['A'].width = 20
    summary_ws.column_dimensions['B'].width = 15
    summary_ws.column_dimensions['C'].width = 40
    summary_ws.column_dimensions['D'].width = 60

    # Create Chart Types Reference sheet
    chart_ws = wb.create_sheet("📈 Chart Types")
    chart_ws.append(["Chart Type", "Description", "Used In"])

    for col in range(1, 4):
        cell = chart_ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    chart_types_info = {
        "radar": "Radar/Spider chart - Multi-dimensional comparison",
        "matrix": "2D matrix - Priority/impact analysis",
        "funnel": "Funnel chart - Sequential conversion stages",
        "hbar": "Horizontal bar - Ranking/comparison",
        "timeline": "Timeline - Sequential events with time",
        "line": "Line chart - Trend over time",
        "multiLine": "Multi-line - Multiple trends comparison",
        "area": "Area chart - Volume over time with emphasis",
        "stacked": "Stacked bar - Component breakdown",
        "groupedColumns": "Grouped columns - Multi-series comparison",
        "donut": "Donut chart - Part-to-whole percentage",
        "pipeline": "Pipeline/Waterfall - Stage conversion",
        "treemap": "Treemap - Hierarchical proportions",
        "heatmap": "Heatmap - Matrix with intensity values",
        "gauge": "Gauge - Single metric progress",
        "modeMap": "Mode map - Strategic options grid",
        "waterfall": "Waterfall - Cumulative effect",
        "segmentFunnel": "Segment funnel - Behavior stage breakdown",
        "table": "Table - Structured data rows"
    }

    # Collect usage
    chart_usage = {chart: [] for chart in chart_types_info.keys()}
    for screen_key, screen_name in screens.items():
        if screen_key not in data:
            continue
        for module in data[screen_key].get('modules', []):
            chart_type = module.get('chart', 'unknown')
            if chart_type in chart_usage:
                chart_usage[chart_type].append(f"{screen_name}: {module.get('title', '')}")

    for chart_type, description in chart_types_info.items():
        used_in = '\n'.join(chart_usage.get(chart_type, []))
        chart_ws.append([chart_type, description, used_in])
        current_row = chart_ws.max_row
        chart_ws.row_dimensions[current_row].height = max(15, len(chart_usage.get(chart_type, [])) * 15)
        for col in range(1, 4):
            cell = chart_ws.cell(row=current_row, column=col)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = thin_border

    chart_ws.column_dimensions['A'].width = 18
    chart_ws.column_dimensions['B'].width = 45
    chart_ws.column_dimensions['C'].width = 70

    # Save file
    output_file = "/Users/khangnhq/Desktop/WorkGalaxy/social-listening/MOCKUP_FULL_FEATURES.xlsx"
    wb.save(output_file)

    print(f"✅ Excel file created: {output_file}")
    print(f"\n📊 Summary:")
    print(f"   • Total screens: {len(screens)}")
    print(f"   • Total modules: {total_modules}")
    print(f"   • Total chart types: {len(chart_types_info)}")

    for screen_key, screen_name in screens.items():
        if screen_key in data:
            module_count = len(data[screen_key].get('modules', []))
            print(f"   • {screen_name}: {module_count} modules")

if __name__ == "__main__":
    html_path = "/Users/khangnhq/Desktop/WorkGalaxy/social-listening/Dotn_v19_hybrid_fixed.html"
    data = extract_data_from_html(html_path)
    create_feature_excel(data)
