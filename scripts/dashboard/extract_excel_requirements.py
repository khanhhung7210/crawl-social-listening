#!/usr/bin/env python3
"""
Extract dashboard requirements from Excel file
"""

import sys
from pathlib import Path
import json

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

try:
    from openpyxl import load_workbook
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
    from openpyxl import load_workbook

def extract_requirements():
    """Extract all dashboard requirements from Excel"""

    excel_path = PROJECT_ROOT / "Dashboard_req_mapping.xlsx"
    wb = load_workbook(excel_path)

    all_requirements = {}

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]

        # Skip if sheet is Gap Analysis or Scoring Logic (not dashboard screens)
        if sheet_name in ['Gap Analysis', 'Scoring Logic Mapping']:
            continue

        print(f"\n{'='*80}")
        print(f"SHEET: {sheet_name}")
        print(f"{'='*80}")

        # Get headers from first row
        headers = [cell.value for cell in sheet[1]]

        dashboards = []

        # Process each row
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            row_data = [cell.value for cell in row]

            # Skip empty rows
            if not any(row_data):
                continue

            # Create dashboard object
            dashboard = {}
            for i, header in enumerate(headers):
                if header and i < len(row_data):
                    dashboard[header] = row_data[i]

            # Check for Vietnamese or English column name
            dashboard_name = dashboard.get('Dashboard') or dashboard.get('Dashboard name')

            if dashboard_name:
                dashboards.append(dashboard)

                # Print summary
                purpose = dashboard.get('Mục đích') or dashboard.get('Purpose')
                formula = dashboard.get('Công thức') or dashboard.get('Formula/Calculation')
                output = dashboard.get('Output/Metrics')

                print(f"\n{row_idx-1}. {dashboard_name}")
                print(f"   Purpose: {purpose}")
                if formula:
                    # Truncate long formulas for display
                    formula_preview = formula[:200] + "..." if len(str(formula)) > 200 else formula
                    print(f"   Formula: {formula_preview}")
                print(f"   Output: {output}")

        all_requirements[sheet_name] = {
            'screen_name': sheet_name,
            'dashboard_count': len(dashboards),
            'dashboards': dashboards
        }

    # Save to JSON
    output_path = PROJECT_ROOT / "data" / "dashboard" / "excel_requirements.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(all_requirements, ensure_ascii=False, indent=2))

    print(f"\n{'='*80}")
    print(f"SUMMARY")
    print(f"{'='*80}")
    for screen, data in all_requirements.items():
        print(f"{screen}: {data['dashboard_count']} dashboards")

    print(f"\n✅ Requirements saved to: {output_path}")

    return all_requirements

if __name__ == "__main__":
    extract_requirements()
