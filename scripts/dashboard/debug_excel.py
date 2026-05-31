#!/usr/bin/env python3
"""
Debug Excel file structure
"""

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]

try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
    from openpyxl import load_workbook

excel_path = PROJECT_ROOT / "Dashboard_req_mapping.xlsx"
wb = load_workbook(excel_path)

print(f"Sheets: {wb.sheetnames}\n")

# Check first sheet in detail
sheet = wb['Overview']

print("=" * 80)
print("OVERVIEW SHEET - First 5 rows")
print("=" * 80)

for row_idx, row in enumerate(sheet.iter_rows(max_row=5), start=1):
    print(f"\nRow {row_idx}:")
    for col_idx, cell in enumerate(row, start=1):
        if cell.value:
            print(f"  Col {col_idx} ({chr(64+col_idx)}): {cell.value}")

print("\n" + "=" * 80)
print("OVERVIEW SHEET - All rows with content")
print("=" * 80)

for row_idx, row in enumerate(sheet.iter_rows(), start=1):
    row_values = [cell.value for cell in row if cell.value]
    if row_values:
        print(f"Row {row_idx}: {len(row_values)} cells with content")
        if row_idx <= 3:  # Show first 3 rows in detail
            print(f"  Values: {row_values}")
