"""
Create comprehensive Excel with ALL features from implementation
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def create_complete_excel():
    """Create comprehensive feature list from implementation"""
    wb = Workbook()
    wb.remove(wb.active)

    # Colors
    header_fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    screen_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    implemented_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    code_ready_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # All modules from implementation
    screens_data = {
        'Overview': {
            'count': 9,
            'modules': [
                {'title': 'Pain Priority Matrix', 'chart': 'matrix', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Revenue Impact Estimate', 'chart': 'revenueImpact', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Lost Customer Signals', 'chart': 'hbar', 'status': '✓ Implemented', 'data': 'branch topics'},
                {'title': "Today's Action Timeline", 'chart': 'timeline', 'status': '✓ Implemented', 'data': 'action_feed'},
                {'title': 'Qualified Signal Summary', 'chart': 'gauge', 'status': '✓ Implemented', 'data': 'platform_summary'},
                {'title': 'Marketing Funnel Leakage', 'chart': 'pipeline', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Channel Signal Quality', 'chart': 'groupedColumns', 'status': '✓ Implemented', 'data': 'platform_summary'},
                {'title': 'Branch Risk Snapshot', 'chart': 'list', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Brand-wide vs Branch-specific Issue Split', 'chart': 'donut', 'status': '✓ Implemented', 'data': 'branch topics'},
            ]
        },
        'Listen': {
            'count': 11,
            'modules': [
                {'title': 'Topic Health Breakdown', 'chart': 'groupedColumns', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Demand Signal Trend', 'chart': 'line', 'status': '✓ Implemented', 'data': 'platform_summary'},
                {'title': 'Positive Theme Treemap', 'chart': 'groupedColumns', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Behavior Segment Funnel', 'chart': 'funnel', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Trend Opportunity Radar', 'chart': 'radar', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Topic Lifecycle Tracker', 'chart': 'timeline', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Signal Source Mix', 'chart': 'donut', 'status': '✓ Implemented', 'data': 'platform_summary'},
                {'title': 'Influencer Signal Watch Lite', 'chart': 'list', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Creative Trigger Board', 'chart': 'hbar', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Topic x Sentiment by Branch', 'chart': 'table', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Priority Action Detail', 'chart': 'gantt', 'status': '✓ Implemented', 'data': 'action_feed'},
            ]
        },
        'Brand Health': {
            'count': 12,
            'modules': [
                {'title': 'Score Waterfall', 'chart': 'waterfall', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Component Trendline', 'chart': 'multiLine', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Fix / Amplify Matrix', 'chart': 'matrix', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Proof of Value Gauge', 'chart': 'gauge', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Category Brand Rank', 'chart': 'hbar', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Campaign Readiness & Benchmark', 'chart': 'radar', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'YMI-style Score Decomposition', 'chart': 'treemap', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Brand Equity & Sentiment Score', 'chart': 'hbar', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Proof of Premium Readiness', 'chart': 'funnel', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Brand Health by Branch', 'chart': 'groupedColumns', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Best Branch Pattern', 'chart': 'list', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Priority Action Detail', 'chart': 'gantt', 'status': '✓ Implemented', 'data': 'action_feed'},
            ]
        },
        'Reputation': {
            'count': 7,
            'modules': [
                {'title': 'Crisis Spike Monitor', 'chart': 'line', 'status': '⚠️ Code Ready', 'data': 'crisis_spike_alerts (empty)'},
                {'title': 'Review Health Heatmap', 'chart': 'table', 'status': '✓ Implemented', 'data': 'branch_intelligence'},
                {'title': 'Response SLA Donut', 'chart': 'donut', 'status': '✓ Implemented', 'data': 'mock data'},
                {'title': 'Social Proof Pipeline', 'chart': 'pipeline', 'status': '✓ Implemented', 'data': 'evidence_cards'},
                {'title': 'Founder / CEO Reputation Watch', 'chart': 'line', 'status': '⚠️ Code Ready', 'data': 'founder_mentions (empty)'},
                {'title': 'Qualified Negative Signal', 'chart': 'stacked', 'status': '⚠️ Code Ready', 'data': 'qualified_negative_summary (empty)'},
                {'title': 'Sensitive Topic Lifecycle', 'chart': 'timeline', 'status': '⚠️ Code Ready', 'data': 'topic_lifecycle (empty)'},
            ]
        },
        'Competitor': {
            'count': 7,
            'modules': [
                {'title': 'Competitive Pressure Radar', 'chart': 'radar', 'status': '✓ Implemented', 'data': 'competitor_pressure'},
                {'title': 'Winning Pattern Comparison', 'chart': 'groupedColumns', 'status': '✓ Implemented', 'data': 'competitor_patterns'},
                {'title': '5-Mode Response Map', 'chart': 'modeMap', 'status': '✓ Implemented', 'data': 'competitor_responses'},
                {'title': 'Competitive Intelligence Readiness', 'chart': 'funnel', 'status': '✓ Implemented', 'data': 'calculated'},
                {'title': 'Competitor Campaign Ranking', 'chart': 'groupedColumns', 'status': '⚠️ Code Ready', 'data': 'campaign_intelligence (empty)'},
                {'title': 'Winning Content Pattern', 'chart': 'hbar', 'status': '⚠️ Code Ready', 'data': 'content_patterns (empty)'},
                {'title': '5W-1H Signal Diagnosis', 'chart': 'modeMap', 'status': '⚠️ Code Ready', 'data': 'signal_diagnosis (empty)'},
            ]
        },
    }

    # Create Summary sheet
    summary_ws = wb.create_sheet("📊 Summary", 0)

    summary_ws.append(["Dashboard Implementation Status - Complete Inventory"])
    summary_ws.merge_cells('A1:E1')
    summary_ws['A1'].fill = header_fill
    summary_ws['A1'].font = Font(bold=True, color="FFFFFF", size=14)
    summary_ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    summary_ws.append([])
    summary_ws.append(["Screen", "Modules Count", "✓ Implemented", "⚠️ Code Ready", "✗ Missing"])

    for col in range(1, 6):
        cell = summary_ws.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_modules = 0
    total_implemented = 0
    total_code_ready = 0

    for screen_name, screen_data in screens_data.items():
        modules = screen_data['modules']
        implemented = len([m for m in modules if '✓' in m['status']])
        code_ready = len([m for m in modules if '⚠️' in m['status']])
        missing = len([m for m in modules if '✗' in m['status']])

        summary_ws.append([screen_name, len(modules), implemented, code_ready, missing])

        total_modules += len(modules)
        total_implemented += implemented
        total_code_ready += code_ready

        current_row = summary_ws.max_row
        for col in range(1, 6):
            cell = summary_ws.cell(row=current_row, column=col)
            cell.border = thin_border

    # Total row
    summary_ws.append([])
    summary_ws.append(["TOTAL", total_modules, total_implemented, total_code_ready, 0])
    total_row = summary_ws.max_row
    for col in range(1, 6):
        cell = summary_ws.cell(row=total_row, column=col)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        cell.border = thin_border

    # Percentages
    summary_ws.append([])
    summary_ws.append(["Status", "Count", "Percentage", "", ""])
    summary_ws.append(["✓ Implemented", total_implemented, f"{total_implemented/total_modules*100:.0f}%", "", ""])
    summary_ws.append(["⚠️ Code Ready", total_code_ready, f"{total_code_ready/total_modules*100:.0f}%", "", ""])
    summary_ws.append(["✗ Missing", 0, "0%", "", ""])

    summary_ws.column_dimensions['A'].width = 20
    summary_ws.column_dimensions['B'].width = 15
    summary_ws.column_dimensions['C'].width = 15
    summary_ws.column_dimensions['D'].width = 15
    summary_ws.column_dimensions['E'].width = 15

    # Create detail sheets for each screen
    for screen_name, screen_data in screens_data.items():
        ws = wb.create_sheet(f"🎯 {screen_name}")

        # Screen header
        ws.merge_cells('A1:E1')
        screen_header = ws['A1']
        screen_header.value = f"{screen_name} - {screen_data['count']} Modules"
        screen_header.fill = screen_fill
        screen_header.font = Font(bold=True, color="FFFFFF", size=13)
        screen_header.alignment = Alignment(horizontal="center", vertical="center")

        # Column headers
        ws.append([])
        headers = ["#", "Module Name", "Chart Type", "Status", "Data Source"]
        ws.append(headers)

        header_row = ws.max_row
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Module data
        for idx, module in enumerate(screen_data['modules'], 1):
            ws.append([
                idx,
                module['title'],
                module['chart'],
                module['status'],
                module['data']
            ])

            current_row = ws.max_row
            for col_num in range(1, 6):
                cell = ws.cell(row=current_row, column=col_num)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            # Color code status
            status_cell = ws.cell(row=current_row, column=4)
            if '✓' in module['status']:
                status_cell.fill = implemented_fill
                status_cell.font = Font(bold=True)
            elif '⚠️' in module['status']:
                status_cell.fill = code_ready_fill
                status_cell.font = Font(bold=True)

        # Column widths
        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 35

    # Save
    output_file = "/Users/khangnhq/Desktop/WorkGalaxy/social-listening/IMPLEMENTATION_COMPLETE_FEATURES.xlsx"
    wb.save(output_file)

    print(f"✅ Complete feature Excel created: {output_file}")
    print(f"\n📊 Summary:")
    print(f"   • Total modules: {total_modules}")
    print(f"   • ✓ Implemented: {total_implemented} ({total_implemented/total_modules*100:.0f}%)")
    print(f"   • ⚠️ Code ready: {total_code_ready} ({total_code_ready/total_modules*100:.0f}%)")
    print(f"\n📋 By Screen:")
    for screen_name, screen_data in screens_data.items():
        print(f"   • {screen_name}: {screen_data['count']} modules")

if __name__ == "__main__":
    create_complete_excel()
