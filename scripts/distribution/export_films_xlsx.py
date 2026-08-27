#!/usr/bin/env python3
"""Export Distribution film catalog + keywords + DB metrics → xlsx."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

def _project_root() -> Path:
    current = Path(__file__).resolve().parent
    for cand in [current, *current.parents]:
        if (cand / "src" / "social_listening").is_dir():
            return cand
    raise RuntimeError(f"Cannot find project root from {__file__}")


PROJECT_ROOT = _project_root()
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from social_listening.paths import DATA_DIR  # noqa: E402

CATALOG = DATA_DIR / "distribution" / "film_catalog.json"
DEFAULT_OUT = (
    PROJECT_ROOT
    / "data"
    / "distribution"
    / "processed"
    / f"distribution_films_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
)


def load_catalog() -> dict:
    return json.loads(CATALOG.read_text(encoding="utf-8"))


def collect_keywords(film: dict) -> list[str]:
    rel = str(film.get("keyword_file") or f"films/{film.get('slug')}.json")
    path = DATA_DIR / "distribution" / rel
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    terms: list[str] = []
    for key in ("keywords", "sub_keywords", "hashtags", "listening_keywords"):
        for value in payload.get(key) or []:
            term = str(value or "").strip()
            if term and term not in terms:
                terms.append(term)
    return terms


def fetch_db_rows() -> tuple[list[tuple], list[tuple]]:
    films_rows: list[tuple] = []
    metrics_rows: list[tuple] = []
    try:
        from social_listening.pg import get_connection
    except Exception as exc:
        print(f"Skip DB sheets: {exc}")
        return films_rows, metrics_rows

    try:
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute(
                """
                SELECT film_slug, film_title, short_title, status, is_active,
                       release_date::text, trailer_date::text, distributor, compare_group
                FROM films
                ORDER BY film_title
                """
            )
            films_rows = cur.fetchall()
            cur.execute(
                """
                SELECT f.film_title, m.metric_date::text, m.platform_code,
                       m.buzz_count, m.post_count, m.comment_count,
                       m.positive_count, m.negative_count, m.neutral_count,
                       m.unique_authors, m.view_count, m.sov_pct
                FROM daily_film_metrics m
                JOIN films f ON f.film_id = m.film_id
                ORDER BY m.metric_date DESC, f.film_title, m.platform_code
                LIMIT 5000
                """
            )
            metrics_rows = cur.fetchall()
    except Exception as exc:
        print(f"DB query failed: {exc}")
    return films_rows, metrics_rows


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    catalog = load_catalog()
    films = catalog.get("films") or []
    db_films, db_metrics = fetch_db_rows()

    wb = Workbook()

    ws = wb.active
    ws.title = "films_catalog"
    ws.append(
        [
            "slug",
            "title",
            "short_title",
            "status",
            "active",
            "distributor",
            "compare_group",
            "aliases",
            "keyword_count",
            "keywords",
        ]
    )
    for film in films:
        kws = collect_keywords(film)
        ws.append(
            [
                film.get("slug"),
                film.get("title"),
                film.get("short_title"),
                film.get("status"),
                bool(film.get("active")),
                film.get("distributor"),
                film.get("compare_group"),
                " | ".join(film.get("aliases") or []),
                len(kws),
                " | ".join(kws),
            ]
        )
    style_header(ws)

    ws2 = wb.create_sheet("films_db")
    ws2.append(
        [
            "film_slug",
            "film_title",
            "short_title",
            "status",
            "is_active",
            "release_date",
            "trailer_date",
            "distributor",
            "compare_group",
        ]
    )
    for row in db_films:
        ws2.append(list(row))
    style_header(ws2)

    ws3 = wb.create_sheet("daily_film_metrics")
    ws3.append(
        [
            "film_title",
            "metric_date",
            "platform",
            "buzz",
            "posts",
            "comments",
            "positive",
            "negative",
            "neutral",
            "unique_authors",
            "views",
            "sov_pct",
        ]
    )
    for row in db_metrics:
        ws3.append(list(row))
    style_header(ws3)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
